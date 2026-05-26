from __future__ import annotations
from pathlib import Path
from typing import List
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from .parser import SlideSpec
from .utils import parse_hex_color, px_to_inches

def _a(v):
    if not v: return None
    v=v.lower().strip()
    return PP_ALIGN.CENTER if v=='center' else (PP_ALIGN.RIGHT if v=='right' else PP_ALIGN.LEFT)

def build_pptx(slides: List[SlideSpec], output_path: Path) -> None:
    prs=Presentation(); prs.slide_width=Inches(px_to_inches(1280)); prs.slide_height=Inches(px_to_inches(720))
    blank=prs.slide_layouts[6]
    for s in slides:
        slide=prs.slides.add_slide(blank)
        for t in s.textboxes:
            sh=slide.shapes.add_textbox(Inches(px_to_inches(t.left_px)),Inches(px_to_inches(t.top_px)),Inches(px_to_inches(t.width_px)),Inches(px_to_inches(t.height_px)))
            p=sh.text_frame.paragraphs[0]; p.text=t.text
            al=_a(t.text_align)
            if al is not None: p.alignment=al
            r=p.runs[0]
            if t.font_size_px is not None: r.font.size=Pt(t.font_size_px)
            if t.font_weight is not None: r.font.bold=t.font_weight>=600
            c=parse_hex_color(t.color)
            if c: r.font.color.rgb=c
        c=s.chart
        data=CategoryChartData(); data.categories=c.chart_info.labels; data.add_series('Series 1', c.chart_info.values)
        cht=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if c.chart_info.index_axis=='y' else XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(px_to_inches(c.left_px)),Inches(px_to_inches(c.top_px)),Inches(px_to_inches(c.width_px)),Inches(px_to_inches(c.height_px)),data).chart
        cht.has_legend=False
        ser=cht.series[0]; col=parse_hex_color(c.chart_info.background_color)
        if col:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb=col
        ser.has_data_labels=True; ser.data_labels.show_value=False; ser.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
        for i,pt in enumerate(ser.points):
            pt.data_label.has_text_frame=True
            v=c.chart_info.values[i]; n=c.chart_info.counts[i]
            pt.data_label.text_frame.text=f"{int(v) if float(v).is_integer() else v}% ({int(n) if float(n).is_integer() else n})"
    output_path.parent.mkdir(parents=True,exist_ok=True); prs.save(str(output_path))
